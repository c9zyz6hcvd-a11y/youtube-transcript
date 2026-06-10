import io
import json
import os
import re
import time
import urllib.request

from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from openai import OpenAI
from flask import Flask, Response, render_template, request, jsonify, stream_with_context, send_file
from youtube_transcript_api import YouTubeTranscriptApi, TranscriptsDisabled, NoTranscriptFound

from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.enums import TA_LEFT
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer

app = Flask(__name__)
yt_api = YouTubeTranscriptApi()

# ─── Register a CJK font for PDF output (covers Traditional + Cantonese) ──────
# Prefer the bundled Noto Sans TC (embeds glyphs → works on any OS, incl. Linux
# servers). Fall back to macOS system fonts for local dev, then Helvetica.
_PDF_FONT, _PDF_FONT_BOLD = "Helvetica", "Helvetica-Bold"
_HERE = os.path.dirname(os.path.abspath(__file__))
_FONT_CANDIDATES = [
    # (regular_path, bold_path, subfont_index)
    (os.path.join(_HERE, "fonts", "NotoSansTC.ttf"),
     os.path.join(_HERE, "fonts", "NotoSansTC-Bold.ttf"), None),
    ("/System/Library/Fonts/STHeiti Light.ttc",
     "/System/Library/Fonts/STHeiti Medium.ttc", 0),
]
for reg_path, bold_path, idx in _FONT_CANDIDATES:
    try:
        if not os.path.exists(reg_path):
            continue
        kw = {"subfontIndex": idx} if idx is not None else {}
        pdfmetrics.registerFont(TTFont("CJK", reg_path, **kw))
        bold_kw = {"subfontIndex": idx} if (idx is not None and os.path.exists(bold_path)) else {}
        pdfmetrics.registerFont(TTFont("CJK-Bold", bold_path if os.path.exists(bold_path) else reg_path, **bold_kw))
        pdfmetrics.registerFontFamily("CJK", normal="CJK", bold="CJK-Bold")
        _PDF_FONT, _PDF_FONT_BOLD = "CJK", "CJK-Bold"
        break
    except Exception:
        continue


def parse_runs(s: str):
    """Split a line into (text, is_bold) runs based on **...** markdown."""
    runs = []
    for part in re.split(r"(\*\*.+?\*\*)", s):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            runs.append((part[2:-2], True))
        else:
            runs.append((part, False))
    return runs


def to_pdf_markup(s: str) -> str:
    """Convert a line to reportlab inline markup, escaping XML and bolding **...**."""
    out = []
    for text, bold in parse_runs(s):
        esc = text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
        out.append(f"<b>{esc}</b>" if bold else esc)
    return "".join(out)

# ─── Language output options ─────────────────────────────────────────────────

LANG_INSTRUCTIONS = {
    "hk": "Output in Hong Kong-style Traditional Chinese. Use natural 口語 style, like a HK blog or LIHKG post. Use HK vocabulary (e.g. 軟件, 網絡, 地鐵, 電單車). Use common HK written Cantonese expressions (e.g. 唔係, 咁樣, 即係, 其實). Avoid overly formal 書面語. Add a 3–5 bullet point summary in Traditional Chinese at the very top.",
    "tw": "Output in Taiwan Traditional Chinese (繁體中文). Use vocabulary and phrasing natural to Taiwan readers. Add a 3–5 bullet point summary in Traditional Chinese at the very top.",
    "cn": "Output in Simplified Chinese (簡體中文). Write in clear, natural Mandarin. Add a 3–5 bullet point summary in Simplified Chinese at the very top.",
    "en": "Output in clear, natural English. Write in a clean, readable style. Add a 3–5 bullet point summary in English at the very top.",
    "ja": "Output in natural Japanese (日本語). Write clearly and naturally. Add a 3–5 bullet point summary in Japanese at the very top.",
    "ko": "Output in natural Korean (한국어). Write clearly and naturally. Add a 3–5 bullet point summary in Korean at the very top.",
}

SUMMARY_LANG_INSTRUCTIONS = {
    "hk": "Write the entire summary in Hong Kong Traditional Chinese (繁體中文). Use clear, natural written Chinese — not overly casual, but readable.",
    "tw": "Write the entire summary in Taiwan Traditional Chinese (繁體中文).",
    "cn": "Write the entire summary in Simplified Chinese (簡體中文).",
    "en": "Write the entire summary in clear, natural English.",
    "ja": "Write the entire summary in Japanese (日本語).",
    "ko": "Write the entire summary in Korean (한국어).",
}

SUMMARY_SYSTEM_PROMPT = """\
You are an expert at extracting and organising information from video transcripts.

Your task: produce a COMPREHENSIVE and DETAILED summary of the transcript below. Your goal is to capture every important piece of information — nothing significant should be missing.

{lang_instruction}

Structure your output as follows:

## 核心重點 / Key Takeaways
5–10 bullet points of the most critical points. Be specific — include numbers, names, facts.

## 詳細內容 / Detailed Notes
Go through the content section by section (use ## headings for each major topic).
For each section:
- Explain the key ideas in full
- Include specific data, statistics, examples, or quotes mentioned
- Do NOT skip over any topic that was discussed for more than a few sentences

## 重要細節 / Important Details
A bullet list of specific facts, figures, names, dates, tools, or resources mentioned that are easy to miss but valuable.

## 結論 / Conclusion
What was the overall message or call to action? What should the viewer take away and do?

Be thorough. It is better to include too much than to miss something important.\
"""

BRIEF_SUMMARY_SYSTEM_PROMPT = """\
You are an expert at summarising video transcripts.

Your task: produce a MEDIUM-LENGTH summary — more than a quick glance, but much shorter than a full detailed breakdown. Aim for something a person can read in 3–5 minutes even for a long podcast.

{lang_instruction}

Structure your output as follows:

## 一句話總結 / TL;DR
One sentence capturing the core message.

## 主要重點 / Key Points
8–12 bullet points. Cover all the significant topics discussed. Be specific — include names, numbers, examples where relevant. Do not merge unrelated points together.

## 各段落重點 / Section Highlights
3–6 short paragraphs (2–4 sentences each), each covering a major theme or segment of the content. Give enough detail that the reader understands what was actually said, not just that the topic was mentioned.

## 結論 / Takeaway
2–3 sentences. What is the main message and what should the reader do or remember?

Be thorough enough to cover a long podcast, but do not go into exhaustive detail on every point.\
"""

BASE_SYSTEM_PROMPT = """\
You are a podcast transcript editor and translator.
Follow these steps:

Detect the source language of the transcript.
Clean up the text — remove filler words (um, ah, 那個, 即係, 咁, 對對對), false starts, repetition, and crosstalk noise.
{lang_instruction}
If translating — translate by meaning and intent, not word-for-word. Preserve the speaker's voice and tone.
Keep proper nouns, brand names, and technical terms in original English if no natural equivalent exists.

Speaker identification — this is important:
- If speaker names are mentioned in the transcript (e.g. "I'm John", "Thanks Sarah"), use their real names.
- If names are not clear, label speakers as 主持人 (Host) and 嘉賓 (Guest), or 嘉賓A / 嘉賓B if there are multiple guests.
- Format each speaker turn on a SINGLE line as: **[Speaker Name]:** followed immediately by their dialogue on the same line.
  Example: **主持人：** 今日我哋傾下健康問題。
- Start a new line with a new speaker label every time the speaker changes.
- Group consecutive sentences from the same speaker together — don't split them unnecessarily.\
"""

# ─── Helpers ──────────────────────────────────────────────────────────────────

def extract_video_id(url: str):
    match = re.search(r"(?:v=|youtu\.be/|embed/|shorts/)([A-Za-z0-9_-]{11})", url)
    return match.group(1) if match else None


def fetch_video_meta(video_id: str) -> dict:
    """Fetch video title via YouTube oEmbed — no API key needed."""
    try:
        oembed = f"https://www.youtube.com/oembed?url=https://www.youtube.com/watch?v={video_id}&format=json"
        req = urllib.request.Request(oembed, headers={"User-Agent": "Mozilla/5.0"})
        with urllib.request.urlopen(req, timeout=5) as resp:
            d = json.loads(resp.read())
            return {
                "title": d.get("title", ""),
                "thumbnail": f"https://img.youtube.com/vi/{video_id}/mqdefault.jpg",
            }
    except Exception:
        return {"title": "", "thumbnail": f"https://img.youtube.com/vi/{video_id}/mqdefault.jpg"}

# ─── Routes ───────────────────────────────────────────────────────────────────

@app.route("/")
def index():
    return render_template("index.html")


@app.route("/transcript", methods=["POST"])
def transcript():
    data = request.get_json()
    url = (data or {}).get("url", "").strip()

    if not url:
        return jsonify({"error": "Please provide a YouTube URL."}), 400

    video_id = extract_video_id(url)
    if not video_id:
        return jsonify({"error": "Could not parse a video ID from that URL."}), 400

    try:
        transcript_list = yt_api.list(video_id)
        all_transcripts = list(transcript_list)
        if not all_transcripts:
            return jsonify({"error": "No transcripts available for this video."}), 400

        # Prefer manually created > auto-generated; prefer English within each group
        def priority(t):
            return (not t.is_generated, t.language_code.lower().startswith("en"))

        meta = sorted(all_transcripts, key=priority, reverse=True)[0]
        snippets = list(yt_api.fetch(video_id, languages=[meta.language_code]))
        plain = " ".join(s.text.replace("\n", " ") for s in snippets)
        video_meta = fetch_video_meta(video_id)

        return jsonify({
            "video_id": video_id,
            "title": video_meta["title"],
            "thumbnail": video_meta["thumbnail"],
            "language": meta.language,
            "plain": plain,
            "entry_count": len(snippets),
        })

    except TranscriptsDisabled:
        return jsonify({"error": "Transcripts are disabled for this video."}), 400
    except NoTranscriptFound:
        return jsonify({"error": "No transcript found for this video."}), 400
    except Exception as e:
        return jsonify({"error": str(e)}), 500


def stream_deepseek(system_prompt: str, user_prompt: str):
    """Shared SSE generator using DeepSeek API."""
    api_key = os.environ.get("DEEPSEEK_API_KEY")
    if not api_key:
        yield f"data: {json.dumps({'error': 'DEEPSEEK_API_KEY is not set.'})}\n\n"
        return

    client = OpenAI(api_key=api_key, base_url="https://api.deepseek.com")

    for attempt in range(3):
        try:
            stream = client.chat.completions.create(
                model="deepseek-v4-flash",
                messages=[
                    {"role": "system", "content": system_prompt},
                    {"role": "user",   "content": user_prompt},
                ],
                stream=True,
                # v4-flash 預設思考mode，transcript 清理唔需要，熄咗佢（快+平）
                extra_body={"thinking": {"type": "disabled"}},
            )
            for chunk in stream:
                text = chunk.choices[0].delta.content or ""
                if text:
                    yield f"data: {json.dumps({'text': text})}\n\n"
            yield "data: [DONE]\n\n"
            return
        except Exception as e:
            error_msg = str(e)
            m = re.search(r'retry in (\d+\.?\d*)s', error_msg)
            if m and attempt < 2:
                wait = min(float(m.group(1)) + 2, 65)
                yield f"data: {json.dumps({'status': f'Rate limited — retrying in {int(wait)}s…'})}\n\n"
                time.sleep(wait)
                continue
            yield f"data: {json.dumps({'error': error_msg})}\n\n"


@app.route("/process", methods=["POST"])
def process():
    data = request.get_json()
    transcript_text = (data or {}).get("transcript", "").strip()
    lang = (data or {}).get("language", "hk")

    if not transcript_text:
        return jsonify({"error": "No transcript provided."}), 400

    lang_instruction = LANG_INSTRUCTIONS.get(lang, LANG_INSTRUCTIONS["hk"])
    system_prompt = BASE_SYSTEM_PROMPT.format(lang_instruction=lang_instruction)
    user_prompt = f"Here is the podcast transcript to process:\n\n{transcript_text}"

    return Response(
        stream_with_context(stream_deepseek(system_prompt, user_prompt)),
        mimetype="text/event-stream",
        headers={"Cache-Control": "no-cache", "X-Accel-Buffering": "no"},
    )


@app.route("/summarise", methods=["POST"])
def summarise():
    data = request.get_json()
    transcript_text = (data or {}).get("transcript", "").strip()
    lang = (data or {}).get("language", "hk")

    if not transcript_text:
        return jsonify({"error": "No transcript provided."}), 400

    lang_instruction = SUMMARY_LANG_INSTRUCTIONS.get(lang, SUMMARY_LANG_INSTRUCTIONS["hk"])
    system_prompt = SUMMARY_SYSTEM_PROMPT.format(lang_instruction=lang_instruction)
    user_prompt = f"Here is the transcript:\n\n{transcript_text}"

    return Response(
        stream_with_context(stream_deepseek(system_prompt, user_prompt)),
        mimetype="text/event-stream",
        headers={"Cache-Control": "no-cache", "X-Accel-Buffering": "no"},
    )


@app.route("/brief", methods=["POST"])
def brief():
    data = request.get_json()
    transcript_text = (data or {}).get("transcript", "").strip()
    lang = (data or {}).get("language", "hk")

    if not transcript_text:
        return jsonify({"error": "No transcript provided."}), 400

    lang_instruction = SUMMARY_LANG_INSTRUCTIONS.get(lang, SUMMARY_LANG_INSTRUCTIONS["hk"])
    system_prompt = BRIEF_SUMMARY_SYSTEM_PROMPT.format(lang_instruction=lang_instruction)
    user_prompt = f"Here is the transcript:\n\n{transcript_text}"

    return Response(
        stream_with_context(stream_deepseek(system_prompt, user_prompt)),
        mimetype="text/event-stream",
        headers={"Cache-Control": "no-cache", "X-Accel-Buffering": "no"},
    )


@app.route("/download/word", methods=["POST"])
def download_word():
    req_data = request.get_json()
    text     = (req_data or {}).get("text", "")
    title    = (req_data or {}).get("title", "Transcript")
    video_id = (req_data or {}).get("video_id", "")

    doc = Document()

    # Narrow margins — use the full page width
    for section in doc.sections:
        section.top_margin    = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin   = Inches(0.9)
        section.right_margin  = Inches(0.9)

    def compact(p, before=0, after=3):
        """Apply tight spacing to a paragraph."""
        p.paragraph_format.space_before  = Pt(before)
        p.paragraph_format.space_after   = Pt(after)
        p.paragraph_format.line_spacing  = Pt(15)

    # Title
    h = doc.add_heading(title, level=1)
    h.alignment = WD_ALIGN_PARAGRAPH.LEFT
    compact(h, before=0, after=2)

    # Source URL
    if video_id:
        p = doc.add_paragraph()
        run = p.add_run(f"https://www.youtube.com/watch?v={video_id}")
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor(0x88, 0x88, 0x88)
        compact(p, after=10)

    # Body
    for line in text.split("\n"):
        s = line.strip()

        if not s or s in ("---", "***", "___"):
            continue

        # Bullet point
        if s.startswith(("* ", "• ", "- ")):
            p = doc.add_paragraph(style="List Bullet")
            for txt, bold in parse_runs(s[2:]):
                p.add_run(txt).bold = bold
            compact(p, after=2)
            continue

        # Heading
        if s.startswith(("## ", "# ")):
            doc.add_heading(s.lstrip("# ").strip(), level=2)
            continue

        # Normal text with inline bold (speaker labels etc.)
        p = doc.add_paragraph()
        for txt, bold in parse_runs(s):
            p.add_run(txt).bold = bold
        compact(p, after=3)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)

    safe = re.sub(r'[^\w\s\-]', '', title)[:50].strip() or "transcript"
    return send_file(
        buf,
        as_attachment=True,
        download_name=f"{safe}.docx",
        mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )


@app.route("/view/pdf", methods=["POST"])
def view_pdf():
    # Accept either form POST (opens inline in a new tab) or JSON
    src      = request.form if request.form else (request.get_json() or {})
    text     = src.get("text", "")
    title    = src.get("title", "Transcript")
    video_id = src.get("video_id", "")

    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=A4,
        topMargin=16 * mm, bottomMargin=16 * mm,
        leftMargin=16 * mm, rightMargin=16 * mm,
        title=title,
    )

    title_style = ParagraphStyle(
        "Title", fontName=_PDF_FONT_BOLD,
        fontSize=16, leading=21, spaceAfter=4, alignment=TA_LEFT, textColor="#111111",
    )
    url_style = ParagraphStyle(
        "Url", fontName=_PDF_FONT, fontSize=8.5, leading=12,
        spaceAfter=12, textColor="#888888",
    )
    body_style = ParagraphStyle(
        "Body", fontName=_PDF_FONT, fontSize=11, leading=17,
        spaceAfter=5, alignment=TA_LEFT, wordWrap="CJK",
    )
    bullet_style = ParagraphStyle(
        "Bullet", parent=body_style, leftIndent=14, bulletIndent=2, spaceAfter=3,
    )
    heading_style = ParagraphStyle(
        "Heading", fontName=_PDF_FONT_BOLD,
        fontSize=13, leading=18, spaceBefore=8, spaceAfter=4, textColor="#111111",
    )

    flow = [Paragraph(to_pdf_markup(title), title_style)]
    if video_id:
        flow.append(Paragraph(f"https://www.youtube.com/watch?v={video_id}", url_style))

    for line in text.split("\n"):
        s = line.strip()
        if not s or s in ("---", "***", "___"):
            flow.append(Spacer(1, 4))
            continue
        if s.startswith(("* ", "• ", "- ")):
            flow.append(Paragraph("• " + to_pdf_markup(s[2:]), bullet_style))
        elif s.startswith(("## ", "# ")):
            flow.append(Paragraph(to_pdf_markup(s.lstrip("# ").strip()), heading_style))
        else:
            flow.append(Paragraph(to_pdf_markup(s), body_style))

    doc.build(flow)
    buf.seek(0)

    safe = re.sub(r'[^\w\s\-]', '', title)[:50].strip() or "transcript"
    # inline → opens in the browser/phone PDF viewer instead of downloading
    return send_file(
        buf,
        mimetype="application/pdf",
        as_attachment=False,
        download_name=f"{safe}.pdf",
    )


if __name__ == "__main__":
    # PORT is provided by the host (Render etc.); default 5055 for local dev.
    port = int(os.environ.get("PORT", 5055))
    debug = os.environ.get("FLASK_DEBUG", "1") == "1"
    app.run(debug=debug, host="0.0.0.0", port=port)
